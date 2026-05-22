# ============================================================
# IMPORTS - ĐÃ DỌN DẸP, XÓA TRÙNG LẶP
# ============================================================
from django.core.paginator import Paginator
from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse, HttpResponse
from django.db.models import Q
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.forms import AuthenticationForm
from django.contrib.auth.models import User
from django.contrib.auth.decorators import login_required
from django.contrib.admin.views.decorators import staff_member_required
from django.views.decorators.http import require_POST
from django.contrib.gis.geos import Point
from django.contrib import messages
from django.core.mail import send_mail, EmailMultiAlternatives
from django.contrib.auth.tokens import default_token_generator
from django.utils.http import urlsafe_base64_encode, urlsafe_base64_decode
from django.utils.encoding import force_bytes, force_str
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from datetime import timedelta, datetime
import csv
import os
import json
import openpyxl
import feedparser
from bs4 import BeautifulSoup
from .models import PhanAnh, HinhAnhPhanAnh, Profile, HoTro, GioiThieu, QuanHuyen
from .forms import DangKyForm, UserEditForm, ProfileEditForm, GioiThieuForm

#trang admin 
@staff_member_required(login_url='login') 
def trang_quan_ly(request):
    from datetime import timedelta
    # 1. Tự động dọn rác quá 7 ngày
    thoi_han = timezone.now() - timedelta(days=7)
    PhanAnh.objects.filter(da_xoa=True, ngay_xoa__lt=thoi_han).delete()

    # --- XỬ LÝ FORM GIỚI THIỆU ---
    obj_gioi_thieu = GioiThieu.objects.first()
    if request.method == 'POST' and 'btn_luu_gioi_thieu' in request.POST:
        form_gt = GioiThieuForm(request.POST, instance=obj_gioi_thieu)
        if form_gt.is_valid():
            form_gt.save()
            messages.success(request, "✅ Đã lưu nội dung Giới thiệu thành công!")
            return redirect('trang_quan_ly') 
    else:
        form_gt = GioiThieuForm(instance=obj_gioi_thieu)

    # 2. PHÂN QUYỀN LÃNH ĐỊA VÀ CHUYÊN MÔN
    user_dang_nhap = request.user

    if user_dang_nhap.is_superuser:
        # 2A. Admin Tổng: Thấy toàn bộ thành phố
        pa_list = PhanAnh.objects.filter(da_xoa=False)
        ds_thung_rac = PhanAnh.objects.filter(da_xoa=True).order_by('-ngay_xoa')
    else:
        # 2B. Nhân viên: Lọc "2 Lớp" (Quận + Chuyên môn)
        try:
            profile = user_dang_nhap.profile
            quan_duoc_giao = profile.quan_quan_ly
            chuyen_mon_giao = profile.chuyen_mon # Ví dụ: 'cap_thoat_nuoc'

            if quan_duoc_giao:
                # 🌟 ĐOẠN ĂN TIỀN: Thêm loai_su_co vào filter
                pa_list = PhanAnh.objects.filter(
                    da_xoa=False, 
                    quan_huyen=quan_duoc_giao,
                    loai_su_co=chuyen_mon_giao # <--- Lọc đúng ngành nghề
                )
                ds_thung_rac = PhanAnh.objects.filter(
                    da_xoa=True, 
                    quan_huyen=quan_duoc_giao,
                    loai_su_co=chuyen_mon_giao # <--- Thùng rác cũng lọc theo ngành
                ).order_by('-ngay_xoa')

                # Cập nhật thông báo cho chuyên nghiệp
                ten_cm = profile.get_chuyen_mon_display()
                messages.info(request, f"Khu vực: {quan_duoc_giao.ten_quan} | Chuyên môn: {ten_cm}")
            else:
                pa_list = PhanAnh.objects.none()
                ds_thung_rac = PhanAnh.objects.none()
                messages.warning(request, "Tài khoản của bạn chưa được phân công khu vực!")
        except Exception as e:
            pa_list = PhanAnh.objects.none()
            ds_thung_rac = PhanAnh.objects.none()

    # --- CÁC PHẦN CÒN LẠI (Sắp xếp, Phân trang...) GIỮ NGUYÊN ---
    pa_list = pa_list.prefetch_related('danh_sach_anh').order_by('-id')
    tat_ca_tieu_de = pa_list.exclude(trang_thai='da_xu_ly').values_list('tieu_de', flat=True).distinct()
    tieu_de_da_chon = request.GET.getlist('filter_tieu_de')
    focus_id = request.GET.get('focus_id')

    if tieu_de_da_chon:
        if focus_id:
            pa_list = pa_list.filter(Q(tieu_de__in=tieu_de_da_chon) | Q(id=focus_id))
        else:
            pa_list = pa_list.filter(tieu_de__in=tieu_de_da_chon)

    ht_list = HoTro.objects.all().order_by('-id')
    paginator_pa = Paginator(pa_list, 15) 
    page_pa = request.GET.get('page_pa') 
    ds_phan_anh = paginator_pa.get_page(page_pa)

    paginator_ht = Paginator(ht_list, 15)
    page_ht = request.GET.get('page_ht')
    ds_ho_tro = paginator_ht.get_page(page_ht)

    filter_role = request.GET.get('filter_role', '')
    ds_user = User.objects.all().order_by('-id')
    if filter_role == 'admin':
        ds_user = ds_user.filter(is_superuser=True)
    elif filter_role == 'staff':
        ds_user = ds_user.filter(is_staff=True, is_superuser=False)
    elif filter_role == 'user':
        ds_user = ds_user.filter(is_staff=False, is_superuser=False)

    ds_quan_huyen = QuanHuyen.objects.all()
    context = {
        'ds_phan_anh': ds_phan_anh,
        'ds_ho_tro': ds_ho_tro,
        'ds_user': ds_user, 
        'ds_thung_rac': ds_thung_rac,
        'tat_ca_tieu_de': tat_ca_tieu_de,
        'tieu_de_da_chon': tieu_de_da_chon,
        'filter_role': filter_role,
        'form_gt': form_gt,
        'ds_quan_huyen': ds_quan_huyen,
    }
    return render(request, 'maps/quan_ly.html', context)

@staff_member_required(login_url='login')
def them_khu_vuc(request):
    if request.user.is_superuser and request.method == 'POST':
        ten_quan_moi = request.POST.get('ten_quan')
        if ten_quan_moi:
            QuanHuyen.objects.get_or_create(ten_quan=ten_quan_moi)
            messages.success(request, f"✅ Đã tạo khu vực mới: {ten_quan_moi}")
    return redirect('trang_quan_ly')

@require_POST
@staff_member_required(login_url='login')
def duyet_phan_anh(request, id):
    # Tìm phản ánh theo ID
    phan_anh = get_object_or_404(PhanAnh, id=id)
    
    # Nếu đang chờ duyệt thì chuyển sang đang xử lý
    if phan_anh.trang_thai == 'cho_duyet':
        phan_anh.trang_thai = 'dang_xu_ly'
    # Nếu đang xử lý thì chuyển sang đã xử lý xong
    elif phan_anh.trang_thai == 'dang_xu_ly':
        phan_anh.trang_thai = 'da_xu_ly'
        
    phan_anh.save()
    return redirect('trang_quan_ly') # Làm xong thì load lại trang quản lý

@require_POST
@staff_member_required(login_url='login')
@login_required
def xoa_phan_anh(request, id):
    # Tìm phản ánh theo ID
    if request.user.is_staff or request.user.is_superuser:
        item = get_object_or_404(PhanAnh, id=id)
    else:
        item = get_object_or_404(PhanAnh, id=id, nguoi_gui=request.user)

    # Thực hiện Xóa mềm
    item.da_xoa = True
    item.ngay_xoa = timezone.now()
    item.save()
    
    messages.success(request, "✅ Đã chuyển vào Thùng rác.")

    # Thông minh: Xóa ở trang nào thì load lại đúng trang đó
    trang_truoc = request.META.get('HTTP_REFERER')
    return redirect(trang_truoc) if trang_truoc else redirect('trang_quan_ly')
@require_POST
@staff_member_required(login_url='login')
def khoi_phuc_phan_anh(request, id):
    item = get_object_or_404(PhanAnh, id=id)
    item.da_xoa = False
    item.ngay_xoa = None
    item.save()
    messages.success(request, "✅ Đã khôi phục phản ánh.")
    return redirect('trang_quan_ly')

@require_POST
@staff_member_required(login_url='login')
def xoa_vinh_vien_phan_anh(request, id):
    item = get_object_or_404(PhanAnh, id=id)
    item.delete()
    messages.warning(request, "🗑️ Đã xóa vĩnh viễn dữ liệu.")
    return redirect('trang_quan_ly')

@staff_member_required(login_url='login')
def them_user(request):
    # CHỈ SUPERUSER MỚI CÓ QUYỀN VÀO ĐÂY
    if not request.user.is_superuser:
        messages.error(request, "⚠️ Chỉ Admin Tổng mới có quyền thêm tài khoản!")
        return redirect('trang_quan_ly')

    if request.method == 'POST':
        # Lấy dữ liệu từ form gửi lên
        username = request.POST.get('username')
        email = request.POST.get('email')
        password = request.POST.get('password')
        role = request.POST.get('role')
        quan_huyen_id = request.POST.get('quan_huyen') 
        
        # --- 🌟 DÒNG MỚI: Lấy CHUYÊN MÔN từ form ---
        chuyen_mon_duoc_chon = request.POST.get('chuyen_mon', 'tong_hop') 

        # 1. Kiểm tra xem tên đăng nhập đã tồn tại chưa
        if User.objects.filter(username=username).exists():
            messages.error(request, f"⚠️ Tên đăng nhập '{username}' đã có người sử dụng!")
            return redirect('trang_quan_ly')
            
        # ==================================================
        # 🚧 BƯỚC CHẶN MỚI: KIỂM TRA TRÙNG EMAIL
        # ==================================================
        if email and User.objects.filter(email=email).exists():
            messages.error(request, f"⚠️ Lỗi: Email '{email}' đã được sử dụng cho một tài khoản khác!")
            return redirect('trang_quan_ly')
        # ==================================================

        try:
            # 2. Tạo User mới (Vượt qua 2 ải trên mới được tới đây)
            new_user = User.objects.create_user(username=username, email=email, password=password)

            # 3. Phân quyền dựa trên lựa chọn
            if role == 'admin':
                new_user.is_staff = True
                new_user.is_superuser = True
                ten_quyen = "Admin Tổng"
            elif role == 'staff':
                new_user.is_staff = True
                new_user.is_superuser = False
                ten_quyen = "Nhân viên"
            else: # user thường
                new_user.is_staff = False
                new_user.is_superuser = False
                ten_quyen = "Người dùng thường"
            
            # Lưu các thay đổi về quyền
            new_user.save()
            
            # 4. KHÚC CODE MỚI ĐÃ ĐƯỢC NÂNG CẤP: XỬ LÝ GẮN QUẬN & CHUYÊN MÔN
            from .models import QuanHuyen, Profile 
            profile, created = Profile.objects.get_or_create(user=new_user)
            
            # Cập nhật chuyên môn (Staff sẽ lấy từ form, Admin/User sẽ lấy mặc định là tong_hop)
            profile.chuyen_mon = chuyen_mon_duoc_chon
            
            if role == 'staff' and quan_huyen_id:
                quan_duoc_chon = QuanHuyen.objects.get(id=quan_huyen_id)
                profile.quan_quan_ly = quan_duoc_chon
                
            # Lưu Profile với các thông tin mới
            profile.save()
            
            messages.success(request, f"✅ Đã tạo thành công tài khoản '{username}' với quyền {ten_quyen}!")
        except Exception as e:
            messages.error(request, f"❌ Đã xảy ra lỗi: {str(e)}")

    return redirect('trang_quan_ly')

@require_POST
@staff_member_required(login_url='login')
def khoa_user(request, id):
    if not request.user.is_superuser:
        messages.error(request, "⚠️ Bạn không có quyền thực hiện thao tác này!")
        return redirect('trang_quan_ly')
    user_can_khoa = get_object_or_404(User, id=id)
    
    # BẢO HIỂM: Không cho phép tự khóa Admin Tổng (chống tự hủy)
    if not user_can_khoa.is_superuser: 
        user_can_khoa.is_active = not user_can_khoa.is_active  # Đảo ngược trạng thái
        user_can_khoa.save()
        
    return redirect('trang_quan_ly')

@require_POST
@staff_member_required(login_url='login')
def xoa_user(request, id):
    if not request.user.is_superuser:
        messages.error(request, "⚠️ Bạn không có quyền thực hiện thao tác này!")
        return redirect('trang_quan_ly')
    user_can_xoa = get_object_or_404(User, id=id)
    
    # BẢO HIỂM: Không cho phép tự xóa Admin Tổng
    if not user_can_xoa.is_superuser:
        user_can_xoa.delete()
        
    return redirect('trang_quan_ly')

# 1. Trang chủ - Lấy số liệu THẬT từ Database
def home(request):
    tong_ho_so = PhanAnh.objects.filter(da_xoa=False).count()
    da_xu_ly = PhanAnh.objects.filter(da_xoa=False, trang_thai='da_xu_ly').count()
    tong_nguoi_dung = User.objects.filter(is_active=True, is_staff=False, is_superuser=False).count()
    ty_le_xu_ly = round((da_xu_ly / tong_ho_so * 100), 1) if tong_ho_so > 0 else 0
    context = {
        'tong_ho_so': tong_ho_so,
        'ty_le_xu_ly': ty_le_xu_ly,
        'tong_nguoi_dung': tong_nguoi_dung,
    }
    return render(request, 'maps/home.html', context)

# 2. Trang Bản đồ
def map_home(request):
    # Lấy dữ liệu điểm cũ ra để hiển thị
    danh_sach = PhanAnh.objects.filter(da_xoa=False)
    return render(request, 'maps/index.html', {'phan_anh': danh_sach})

# 3. Trang Danh sách
def list_view(request):
    danh_sach_all = PhanAnh.objects.filter(da_xoa=False).order_by('-id')
    paginator = Paginator(danh_sach_all, 15)
    page_number = request.GET.get('page')
    phan_anh = paginator.get_page(page_number)
    return render(request, 'maps/list.html', {'phan_anh': phan_anh})

def register_view(request):
    if request.method == 'POST':
        form = DangKyForm(request.POST) 
        
        if form.is_valid():
            # Lấy email từ form ra để check trùng
            email = form.cleaned_data.get('email')
            
            # 1. KIỂM TRA TRÙNG EMAIL
            if User.objects.filter(email=email).exists():
                messages.error(request, 'Email này đã được sử dụng. Vui lòng chọn Email khác!')
                return render(request, 'maps/register.html', {'form': form})

            # 2. LƯU USER NHƯNG KHÓA LẠI (commit=False để khoan lưu vội)
            user = form.save(commit=False)
            user.is_active = False # Bắt buộc kích hoạt mail mới cho xài
            user.save()

            # 3. TẠO MÃ BÍ MẬT & LINK KÍCH HOẠT
            uid = urlsafe_base64_encode(force_bytes(user.pk))
            token = default_token_generator.make_token(user)
            domain = request.get_host()
            link_kich_hoat = f"http://{domain}{reverse('kich_hoat_tai_khoan', kwargs={'uidb64': uid, 'token': token})}"

            # 4. BẮN MAIL QUA MAILTRAP
            subject = 'Kích hoạt tài khoản SafeCity'
            html_content = f"""
                <div style="font-family: Arial, sans-serif; padding: 20px; background-color: #f4f6f8;">
                    <div style="max-width: 600px; margin: 0 auto; background: white; padding: 30px; border-radius: 10px;">
                        <h2 style="color: #0d6efd;">Chào mừng đến với SafeCity!</h2>
                        <p>Bạn vừa tạo một tài khoản với tên đăng nhập là: <b>{user.username}</b></p>
                        <p>Vui lòng nhấn vào nút bên dưới để xác thực email và kích hoạt tài khoản của bạn:</p>
                        <div style="text-align: center; margin: 30px 0;">
                            <a href="{link_kich_hoat}" style="background-color: #28a745; color: white; padding: 12px 25px; text-decoration: none; border-radius: 5px; font-weight: bold;">Xác Nhận Email</a>
                        </div>
                    </div>
                </div>
            """
            email_msg = EmailMultiAlternatives(subject, "Vui lòng bật HTML để xem mail", 'admin@safecity.com', [email])
            email_msg.attach_alternative(html_content, "text/html")
            email_msg.send()

            # XÓA DÒNG login() CŨ, THAY BẰNG THÔNG BÁO VÀ CHUYỂN HƯỚNG
            messages.success(request, 'Đăng ký thành công! Vui lòng kiểm tra Email (Mailtrap) để kích hoạt tài khoản.')
            return redirect('login') # Chuyển về trang đăng nhập

    else:
        form = DangKyForm()
        
    return render(request, 'maps/register.html', {'form': form})

def kich_hoat_tai_khoan(request, uidb64, token):
    try:
        # Giải mã xem ai đang bấm link
        uid = force_str(urlsafe_base64_decode(uidb64))
        user = User.objects.get(pk=uid)
    except (TypeError, ValueError, OverflowError, User.DoesNotExist):
        user = None

    # Kiểm tra mã Token có đúng không
    if user is not None and default_token_generator.check_token(user, token):
        user.is_active = True # MỞ KHÓA THÀNH CÔNG
        user.save()
        messages.success(request, 'Tài khoản của bạn đã được kích hoạt thành công! Giờ có thể đăng nhập.')
        return redirect('login') # Đổi chữ 'login' cho khớp với tên trang đăng nhập của ông
    else:
        messages.error(request, 'Link kích hoạt không hợp lệ hoặc đã hết hạn!')
        return redirect('register') # Đổi chữ 'register' cho khớp với tên trang đăng ký

# 5. ĐĂNG NHẬP
def login_view(request):
    if request.method == 'POST':
        form = AuthenticationForm(data=request.POST)
        if form.is_valid():
            user = form.get_user()
            login(request, user)
            if 'next' in request.GET:
                return redirect(request.GET.get('next'))
            return redirect('map_home')
    else:
        form = AuthenticationForm()
    return render(request, 'maps/login.html', {'form': form})

# 6. ĐĂNG XUẤT
def logout_view(request):
    logout(request)
    return redirect('login')

# 7. API Lưu phản ánh (ĐÃ SỬA: LƯU NGƯỜI GỬI)
def luu_phan_anh(request):
    if request.method == 'POST':
        try:
            # --- 1. LẤY DỮ LIỆU TỪ FORM ---
            # Lấy chuỗi gộp dạng 'cap_thoat_nuoc|Ngập nước'
            raw_data = request.POST.get('du_lieu_su_co')
            
            if raw_data and "|" in raw_data:
                # Tách đôi: ma_chuyen_mon = 'cap_thoat_nuoc', ten_su_co = 'Ngập nước'
                ma_chuyen_mon, ten_su_co = raw_data.split('|', 1)
            else:
                ma_chuyen_mon = 'khac'
                ten_su_co = 'Sự cố không xác định'

            mota = request.POST.get('mo_ta', '')
            toado = request.POST.get('points_data', '[]')
            diachi = request.POST.get('dia_chi', '') 
            
            # --- KIỂM TRA GIỚI HẠN GỬI BÀI (SPAM CHECK) ---
            if request.user.is_authenticated:
                mot_gio_truoc = timezone.now() - timedelta(hours=1)
                so_bai_da_gui = PhanAnh.objects.filter(
                    nguoi_gui=request.user,
                    thoi_gian__gte=mot_gio_truoc
                ).count()

                if so_bai_da_gui >= 3:
                    return JsonResponse({
                        'success': False,
                        'message': '⏳ Bạn đã đạt giới hạn gửi 3 phản ánh trong 1 giờ. Vui lòng nghỉ ngơi và thử lại sau nhé!'
                    })
            
            # --- 2. KHỞI TẠO ĐỐI TƯỢNG PHẢN ÁNH ---
            # Sử dụng ten_su_co làm tiêu đề và ma_chuyen_mon làm phân loại chuyên môn
            new_pa = PhanAnh(
                tieu_de=ten_su_co,      # Hiển thị: "Ngập nước"
                loai_su_co=ma_chuyen_mon, # Lưu mã: "cap_thoat_nuoc" để lọc việc
                mo_ta=mota,
                dia_chi=diachi,
                du_lieu_toa_do=toado,
            )

            # --- KIỂM TRA TRÙNG ẢNH ---
            hinh_anh_list = request.FILES.getlist('hinh_anh')
            for f in hinh_anh_list:
                ten_goc, duoi_file = os.path.splitext(f.name)
                ten_goc_sach = ten_goc.replace(" ", "_")
                chuoi_tim_kiem = f"/{ten_goc_sach}"
                
                trung_anh_chinh = PhanAnh.objects.filter(hinh_anh__icontains=chuoi_tim_kiem).exists()
                trung_anh_phu = HinhAnhPhanAnh.objects.filter(hinh_anh__icontains=chuoi_tim_kiem).exists()

                if trung_anh_chinh or trung_anh_phu:
                    return JsonResponse({
                        'success': False,
                        'message': f"Tên ảnh '{f.name}' đã bị trùng trên hệ thống. Vui lòng đổi tên file khác!"
                    })

            # --- 3. XỬ LÝ TỌA ĐỘ (POSTGIS) VÀ AUTO-ASSIGNMENT ---
            if toado and toado != "[]":
                try:
                    points_list = json.loads(toado)
                    if len(points_list) > 0:
                        lat = float(points_list[0]['lat'])
                        lng = float(points_list[0]['lng'])
                        
                        diem_su_co = Point(lng, lat, srid=4326)
                        new_pa.vi_tri = diem_su_co
                        
                        # Ma thuật tự động tìm Quận/Huyện dựa trên tọa độ
                        khu_vuc_chua = QuanHuyen.objects.filter(ranh_gioi__contains=diem_su_co).first()
                        
                        if khu_vuc_chua:
                            new_pa.quan_huyen = khu_vuc_chua
                        else:
                            return JsonResponse({
                                'success': False,
                                'message': '🚫 Vị trí này nằm ngoài phạm vi TP.HCM! Vui lòng chọn lại.'
                            })
                except Exception as e:
                    print("⚠️ Lỗi PostGIS:", e)

            # Gắn người gửi
            if request.user.is_authenticated:
                new_pa.nguoi_gui = request.user
            
            # Lưu lần 1
            new_pa.save()

            # --- 4. XỬ LÝ LƯU NHIỀU HÌNH ẢNH ---
            if hinh_anh_list:
                new_pa.hinh_anh = hinh_anh_list[0]
                new_pa.save() 
                
                for file_anh in hinh_anh_list[1:]:
                    HinhAnhPhanAnh.objects.create(
                        phan_anh=new_pa, 
                        hinh_anh=file_anh
                    )

            return JsonResponse({'success': True, 'message': 'Gửi phản ánh thành công!'})
        
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})

    return JsonResponse({'success': False, 'message': 'Yêu cầu không hợp lệ!'})
# 8. TRANG HỒ SƠ CÁ NHÂN (ĐÃ SỬA: LỌC ĐÚNG CỘT nguoi_gui)
@login_required
def profile(request):
    # Dùng .filter(nguoi_gui=...) thay vì user=...
    danh_sach = PhanAnh.objects.filter(nguoi_gui=request.user).order_by('-id')
    
    context = {
        'user': request.user,
        'danh_sach': danh_sach,
        'tong_so': danh_sach.count()
    }
    return render(request, 'maps/profile.html', context)

@login_required
def chi_tiet_ho_so(request, id_ho_so):
    # Lấy hồ sơ kèm theo toàn bộ ảnh trong kho phụ 'danh_sach_anh'
    ho_so = get_object_or_404(PhanAnh.objects.prefetch_related('danh_sach_anh'), id=id_ho_so)
    
    # Phân quyền: Staff/Superuser xem được tất cả
    # User thường chỉ xem được phản ánh của chính mình
    if not (request.user.is_staff or request.user.is_superuser):
        if ho_so.nguoi_gui != request.user:
            from django.core.exceptions import PermissionDenied
            raise PermissionDenied
    
    return render(request, 'maps/detail.html', {'ho_so': ho_so})


@login_required
def edit_profile(request):
    # Đảm bảo user luôn có profile (tránh lỗi cho user cũ)
    Profile.objects.get_or_create(user=request.user)

    if request.method == 'POST':
        user_form = UserEditForm(request.POST, instance=request.user)
        profile_form = ProfileEditForm(request.POST, request.FILES, instance=request.user.profile)
        
        if user_form.is_valid() and profile_form.is_valid():
            user_form.save()
            profile_form.save()
            return redirect('profile')
    else:
        user_form = UserEditForm(instance=request.user)
        profile_form = ProfileEditForm(instance=request.user.profile)
    
    context = {
        'user_form': user_form,
        'profile_form': profile_form
    }
    return render(request, 'maps/edit_profile.html', context)
# Bổ sung các hàm bị thiếu
def trang_thong_ke(request):
    from django.contrib.auth.models import User
    tong_so = PhanAnh.objects.count()
    cho_duyet = PhanAnh.objects.filter(trang_thai='cho_duyet').count()
    dang_xu_ly = PhanAnh.objects.filter(trang_thai='dang_xu_ly').count()
    da_xu_ly = PhanAnh.objects.filter(trang_thai='da_xu_ly').count()
    ty_le = 0
    if tong_so > 0: ty_le = round((da_xu_ly / tong_so) * 100, 1)

    context = {'tong_so': tong_so, 'cho_duyet': cho_duyet, 'dang_xu_ly': dang_xu_ly, 'da_xu_ly': da_xu_ly, 'ty_le': ty_le, 'so_nguoi_dung': User.objects.count(), 'bai_moi': PhanAnh.objects.all().order_by('-id')[:5]}
    return render(request, 'maps/thong_ke.html', context)

@login_required
def cap_nhat_trang_thai(request, id_phan_anh, trang_thai_moi):
    phan_anh = get_object_or_404(PhanAnh, id=id_phan_anh)
    phan_anh.trang_thai = trang_thai_moi
    phan_anh.save()
    return redirect('quan_ly_hien_truong')

def api_get_points(request):
    user_dang_nhap = request.user
    
    # 1. BỘ LỌC PHÂN QUYỀN (Giữ nguyên logic cực xịn)
    if user_dang_nhap.is_superuser:
        danh_sach = PhanAnh.objects.filter(da_xoa=False)
    else:
        try:
            quan_duoc_giao = user_dang_nhap.profile.quan_quan_ly
            if quan_duoc_giao:
                danh_sach = PhanAnh.objects.filter(da_xoa=False, quan_huyen=quan_duoc_giao)
            else:
                danh_sach = PhanAnh.objects.none()
        except:
            danh_sach = PhanAnh.objects.none()

    # 2. XUẤT DỮ LIỆU RA JSON (Sử dụng trực tiếp sức mạnh PostGIS)
    data = []
    for item in danh_sach:
        try:
            # Chỉ lấy những bài đã được lưu tọa độ PostGIS thành công
            if item.vi_tri:
                lat = item.vi_tri.y # PostGIS y là Vĩ độ (Latitude)
                lng = item.vi_tri.x # PostGIS x là Kinh độ (Longitude)
                
                # Xử lý ảnh cực kỳ an toàn, bao không crash hệ thống
                img_url = ""
                if item.hinh_anh and hasattr(item.hinh_anh, 'url'):
                    try:
                        img_url = item.hinh_anh.url
                    except ValueError:
                        pass # Nếu rỗng thì bỏ qua, xài chuỗi rỗng
                        
                data.append({
                    'id': item.id, 
                    'title': item.tieu_de, 
                    'lat': lat, 
                    'lng': lng, 
                    'status': item.trang_thai, 
                    'image_url': img_url, 
                    'detail_url': f"/chi-tiet/{item.id}/"
                })
        except Exception as e:
            print(f"⚠️ Lỗi xuất dữ liệu ở bài #{item.id}: {e}")
            
    return JsonResponse(data, safe=False)


@login_required
def quan_ly_hien_truong(request):
    # Lấy các điểm ĐANG XỬ LÝ (đang thi công)
    danh_sach = PhanAnh.objects.filter(trang_thai='dang_xu_ly').order_by('-thoi_gian')
    
    context = {
        'danh_sach': danh_sach,
        'so_luong': danh_sach.count()
    }
    return render(request, 'maps/quan_ly_hien_truong.html', context)

@staff_member_required(login_url='login')
def export_excel(request):
    # 1. Cấu hình response là file CSV
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="bao_cao_su_co.csv"'
    
    # 2. Fix lỗi font tiếng Việt (BOM header)
    response.write(u'\ufeff'.encode('utf8')) 

    # 3. Tạo bút ghi
    writer = csv.writer(response)
    
    # 4. Ghi dòng tiêu đề
    writer.writerow(['ID', 'Tiêu đề', 'Tọa độ', 'Thời gian', 'Trạng thái'])

    # 5. Lấy dữ liệu và ghi từng dòng
    for pa in PhanAnh.objects.all().order_by('-id'):
        writer.writerow([
            pa.id, 
            pa.tieu_de, 
            pa.du_lieu_toa_do, 
            pa.thoi_gian.strftime("%d/%m/%Y %H:%M"), 
            pa.get_trang_thai_display()
        ])

    return response


def tin_tuc(request):
    # Link RSS VnExpress
    rss_url = "https://vnexpress.net/rss/thoi-su.rss"
    feed = feedparser.parse(rss_url)
    
    # Bộ lọc từ khóa (Giữ nguyên)
    keywords = [
        'giao thông', 'đô thị', 'ngập', 'kẹt xe', 'ùn tắc', 
        'cầu', 'đường', 'hầm', 'chung cư', 'quy hoạch', 
        'môi trường', 'rác', 'metro', 'xe buýt', 'vỉa hè',
        'cây xanh', 'chiếu sáng', 'hcm', 'sài gòn', 'quận'
    ]

    posts = []
    for entry in feed.entries:
        content_to_check = (entry.title + entry.description).lower()
        
        if any(word in content_to_check for word in keywords):
            # XỬ LÝ ẢNH BẰNG BEAUTIFULSOUP (MỚI & XỊN) 
            image_url = 'https://s1.vnecdn.net/vnexpress/restruct/i/v884/logo_default.jpg' # Hình mặc định
            
            # Dùng "cái muỗng" để múc nội dung HTML trong phần mô tả
            soup = BeautifulSoup(entry.description, 'lxml')
            
            # Tìm thẻ <img> đầu tiên
            img_tag = soup.find('img')
            
            # Nếu tìm thấy thẻ img và nó có thuộc tính src (link ảnh)
            if img_tag and img_tag.get('src'):
                image_url = img_tag.get('src')

            # XỬ LÝ MÔ TẢ CHO SẠCH SẼ 
            # Dùng soup.get_text() để lấy toàn bộ chữ, bỏ hết các thẻ HTML thừa
            summary_text = soup.get_text()
            
            # Cắt bỏ mấy cái chữ thừa ở cuối nếu có (ví dụ ">> Chi tiết")
            if '>>' in summary_text:
                summary_text = summary_text.split('>>')[0]

            posts.append({
                'title': entry.title,
                'link': entry.link,
                'published': entry.published,
                'summary': summary_text.strip(), # Xóa khoảng trắng thừa đầu đuôi
                'image': image_url
            })

        if len(posts) >= 12:
            break

    return render(request, 'maps/tin_tuc.html', {'news_list': posts})


def huong_dan(request):
    return render(request, 'maps/huong_dan.html')



def hotline(request):
    return render(request, 'maps/hotline.html')

def cskh(request):
    if request.method == 'POST':
        # 1. Lấy dữ liệu từ form người dùng gửi lên
        ho_ten = request.POST.get('ho_ten')
        email = request.POST.get('email')
        sdt = request.POST.get('sdt')
        chu_de = request.POST.get('chu_de')
        noi_dung = request.POST.get('noi_dung')

        # 2. Tạo phiếu hỗ trợ mới
        phieu = HoTro(
            ho_ten=ho_ten,
            email=email,
            sdt=sdt,
            chu_de=chu_de,
            noi_dung=noi_dung
        )
        
        # Nếu người dùng đang đăng nhập, gắn luôn user vào để dễ theo dõi
        if request.user.is_authenticated:
            phieu.nguoi_gui = request.user
            
        phieu.save()

        # 3. Thông báo thành công
        messages.success(request, "Đã gửi yêu cầu hỗ trợ! Chúng tôi sẽ phản hồi qua Email sớm nhất.")
        return redirect('cskh') # Load lại trang để xóa form

    return render(request, 'maps/cskh.html')
from django.contrib.gis.measure import D

def api_quet_vung_postgis(request):
    """
    TOOL GIS BACKEND: Tìm các sự cố nằm trong vòng tròn bán kính R 
    (Sử dụng hàm __distance_lte của PostGIS)
    """
    # 1. Nhận tọa độ tâm và bán kính từ Frontend gửi lên (mặc định lấy chợ Bến Thành)
    lat = float(request.GET.get('lat', 10.7725))
    lng = float(request.GET.get('lng', 106.6980))
    ban_kinh_met = float(request.GET.get('radius', 500)) 

    # 2. Tạo điểm tâm (Nhớ là Kinh độ trước, Vĩ độ sau)
    tam_diem = Point(lng, lat, srid=4326)

    # 3. QUAN TRỌNG NHẤT: Truy vấn DB bằng hàm không gian của GeoDjango/PostGIS
    ds_su_co = PhanAnh.objects.filter(
        vi_tri__distance_lte=(tam_diem, D(m=ban_kinh_met))
    )

    # 4. Trả kết quả ra dạng JSON
    data = []
    for sc in ds_su_co:
        data.append({
            'id': sc.id,
            'tieu_de': sc.tieu_de,
            # Tính luôn khoảng cách thực tế từ tâm đến điểm đó để show ra báo cáo
            'khoang_cach': round(sc.vi_tri.distance(tam_diem) * 100000, 2) # Nhân để đổi ra mét tương đối
        })

    return JsonResponse({'tam_diem': {'lat': lat, 'lng': lng}, 'tong_so': len(data), 'du_lieu': data})

@staff_member_required(login_url='login')
def tra_loi_ho_tro(request, id):
    # 1. Tìm cái yêu cầu hỗ trợ theo ID
    ht = get_object_or_404(HoTro, id=id)

    if request.method == 'POST':
        # 2. Lấy nội dung Admin nhập từ form
        noidung_admin_tra_loi = request.POST.get('noidung_phan_hoi')

        # 3. GÓI GHÉM VÀ GỬI EMAIL BẰNG MAILTRAP
        tieu_de = f"[Urban Manager] Phản hồi yêu cầu hỗ trợ #{ht.id}"
        loi_nhan = f"""Chào {ht.ho_ten},

Chúng tôi đã nhận được yêu cầu của bạn về vấn đề: {ht.get_chu_de_display()}.
Nội dung bạn gửi: "{ht.noi_dung}"

PHẢN HỒI TỪ BAN QUẢN TRỊ:
{noidung_admin_tra_loi}

Trân trọng,
Đội ngũ Urban Manager.
"""
        # Bấm nút gửi!
        send_mail(
            subject=tieu_de,
            message=loi_nhan,
            from_email='admin@urbanmanager.com', 
            recipient_list=[ht.email], 
            fail_silently=False,
        )

        # 4. Cập nhật lại Database
        ht.da_xu_ly = True
        ht.phan_hoi_admin = noidung_admin_tra_loi
        ht.save()

        # Báo cáo thành công và quay về trang quản lý của ông (tên là trang_quan_ly)
        messages.success(request, f'Đã gửi email trả lời cho {ht.ho_ten} thành công!')
        return redirect('trang_quan_ly') 

    # Nếu chưa bấm gửi thì gọi cái giao diện soạn thư ra
    return render(request, 'maps/tra_loi_ho_tro.html', {'ht': ht})

def gioi_thieu_view(request):
    # Lấy bản ghi giới thiệu mới nhất
    info = GioiThieu.objects.first() 
    return render(request, 'maps/gioi_thieu.html', {'info': info})

# API MỚI: TRẢ VỀ DỮ LIỆU RANH GIỚI POLYGON CHO BẢN ĐỒ ADMIN
@staff_member_required(login_url='login')
def api_get_ranh_gioi(request):
    user_dang_nhap = request.user
    
    # 1. Lọc danh sách quận theo quyền hạn
    if user_dang_nhap.is_superuser:
        ds_quan = QuanHuyen.objects.all()
    else:
        try:
            quan_duoc_giao = user_dang_nhap.profile.quan_quan_ly
            ds_quan = QuanHuyen.objects.filter(id=quan_duoc_giao.id) if quan_duoc_giao else QuanHuyen.objects.none()
        except:
            ds_quan = QuanHuyen.objects.none()

    # 2. BUILD GEOJSON KÈM THỐNG KÊ
    features = []
    for q in ds_quan:
        if q.ranh_gioi:
            # --- ĐOẠN TÍNH TOÁN SỐ LIỆU ---
            tong_so = PhanAnh.objects.filter(quan_huyen=q, da_xoa=False).count()
            cho_duyet = PhanAnh.objects.filter(quan_huyen=q, da_xoa=False, trang_thai='cho_duyet').count()
            da_xong = PhanAnh.objects.filter(quan_huyen=q, da_xoa=False, trang_thai='da_xu_ly').count()

            # --- TÍNH NĂNG MỚI: TÌM NGƯỜI QUẢN LÝ ---
            # Lọc ra những User có profile gán với quận này
            nguoi_quan_ly = User.objects.filter(profile__quan_quan_ly=q, is_active=True).values_list('username', flat=True)
            if nguoi_quan_ly:
                chuoi_quan_ly = ", ".join(nguoi_quan_ly) # Nếu có 2 người cùng quản lý 1 quận thì nối tên lại
            else:
                chuoi_quan_ly = "Chưa phân công"

            features.append({
                "type": "Feature",
                "geometry": json.loads(q.ranh_gioi.geojson), 
                "properties": {
                    "id": q.id,
                    "ten_quan": q.ten_quan,
                    "mau_sac": q.mau_sac,
                    "nguoi_quan_ly": chuoi_quan_ly, # Gửi thêm cái tên này ra cho Frontend
                    "thong_ke": {
                        "tong": tong_so,
                        "cho_duyet": cho_duyet,
                        "hoan_thanh": da_xong
                    }
                }
            })
            
    geojson_dict = {
        "type": "FeatureCollection",
        "features": features
    }
    
    return HttpResponse(json.dumps(geojson_dict), content_type='application/json')
@require_POST
@staff_member_required(login_url='login')
def xoa_tat_ca_thung_rac(request):
    user_dang_nhap = request.user
    
    if user_dang_nhap.is_superuser:
        # 1. Admin Tổng: Xóa toàn bộ thùng rác của toàn thành phố
        PhanAnh.objects.filter(da_xoa=True).delete()
        messages.success(request, "🗑️ Đã dọn sạch vĩnh viễn TOÀN BỘ thùng rác trên hệ thống!")
    else:
        # 2. Nhân viên: Chỉ dọn sạch thùng rác thuộc quận/huyện mình quản lý
        try:
            quan_duoc_giao = user_dang_nhap.profile.quan_quan_ly
            if quan_duoc_giao:
                PhanAnh.objects.filter(da_xoa=True, quan_huyen=quan_duoc_giao).delete()
                messages.success(request, f"🗑️ Đã dọn sạch thùng rác thuộc {quan_duoc_giao.ten_quan}!")
            else:
                messages.error(request, "Lỗi: Bạn chưa được phân công khu vực quản lý.")
        except:
            messages.error(request, "Lỗi: Không xác định được hồ sơ người dùng.")
            
    return redirect('trang_quan_ly')

# --- HÀM 1: XUẤT EXCEL ---
@staff_member_required(login_url='login')
def xuat_excel_lich_su(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lich_Su_Xu_Ly"
    
    # Lấy ngày giờ thực tế
    bay_gio = datetime.now()
    ten_file = f"bao_cao_phan_tich_{bay_gio.strftime('%d_%m_%Y')}.xlsx"
    
    # 1. Tạo hàng tiêu đề
    headers = ['ID', 'Người gửi', 'Tiêu đề', 'Địa chỉ', 'Thời gian gửi', 'Trạng thái']
    ws.append(headers)

    MA_TRANG_THAI = 'da_xu_ly' 
    
    user_dang_nhap = request.user
    
    if user_dang_nhap.is_superuser:
        # Admin Tổng: Lấy hết toàn bộ thành phố
        lich_su = PhanAnh.objects.filter(trang_thai=MA_TRANG_THAI, da_xoa=False).order_by('-id')
    else:
        # Nhân viên: Chỉ lấy đúng Quận và Chuyên môn
        profile = user_dang_nhap.profile
        lich_su = PhanAnh.objects.filter(
            trang_thai=MA_TRANG_THAI, 
            da_xoa=False,
            quan_huyen=profile.quan_quan_ly,
            loai_su_co=profile.chuyen_mon
        ).order_by('-id')

    # ==========================================

    # 2. Đổ dữ liệu vào file Excel
    for pa in lich_su:
        thoi_gian_str = pa.thoi_gian.strftime('%d/%m/%Y %H:%M') if pa.thoi_gian else ""
        ws.append([
            f"#{pa.id}", 
            pa.nguoi_gui.username if pa.nguoi_gui else "Ẩn danh",
            pa.tieu_de,
            pa.dia_chi,
            thoi_gian_str,
            "Đã xử lý xong"
        ])

    # 3. Thêm khu vực chữ ký ở góc dưới
    last_row = ws.max_row + 3
    profile = request.user.profile
    
    ws.cell(row=last_row, column=4, value=f"Ngày xuất: {bay_gio.strftime('%d/%m/%Y %H:%M')}")
    ten_quan = profile.quan_quan_ly.ten_quan if profile.quan_quan_ly else "Toàn thành phố"
    ws.cell(row=last_row + 1, column=4, value=f"Khu vực quản lý: {ten_quan}")
    ws.cell(row=last_row + 2, column=4, value=f"Chuyên môn: {profile.get_chuyen_mon_display()}")
    ws.cell(row=last_row + 4, column=4, value="Người lập biểu")
    ws.cell(row=last_row + 5, column=4, value="(Ký và ghi rõ họ tên)")
    ws.cell(row=last_row + 7, column=4, value=request.user.get_full_name() or request.user.username)

    # 4. Trả về file tải xuống
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{ten_file}"'
    wb.save(response)
    return response

# --- HÀM 2: XÓA TẤT CẢ LỊCH SỬ ---
@require_POST
@staff_member_required(login_url='login')
def xoa_tat_ca_lich_su(request):
    # Thay vì xóa vĩnh viễn, mình đổi trạng thái da_xoa=True để ném vào thùng rác cho an toàn
    PhanAnh.objects.filter(trang_thai='da_xu_ly', da_xoa=False).update(da_xoa=True)
    messages.success(request, "Đã dọn dẹp sạch sẽ toàn bộ lịch sử!")
    return redirect('trang_quan_ly')